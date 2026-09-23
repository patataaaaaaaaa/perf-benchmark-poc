from __future__ import annotations

import argparse
import json
import os
import tracemalloc
from pathlib import Path


MIB = 1024 * 1024


def prepare_openpyxl(path: Path, cells: int) -> dict[str, int]:
    from openpyxl import Workbook

    columns = 10
    if cells <= 0 or cells % columns:
        raise ValueError("openpyxl cell count must be positive and divisible by 10")
    rows = cells // columns
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("data")
    for row in range(rows):
        worksheet.append([row * columns + column for column in range(columns)])
    workbook.save(path)
    return {"rows": rows, "columns": columns, "cells": cells}


def prepare_fsspec(path: Path, size_mib: int) -> dict[str, int]:
    if size_mib <= 0:
        raise ValueError("fsspec size must be positive")
    chunk = b"x" * MIB
    with path.open("wb") as stream:
        for _ in range(size_mib):
            stream.write(chunk)
        stream.flush()
        os.fsync(stream.fileno())
    return {"size_mib": size_mib, "bytes": size_mib * MIB}


def prepare_cachetools(path: Path, size_mib: int) -> dict[str, int]:
    if size_mib <= 0:
        raise ValueError("cachetools size must be positive")
    details = {"size_mib": size_mib, "bytes": size_mib * MIB}
    path.write_text(json.dumps(details) + "\n", encoding="utf-8")
    return details


def drop_file_cache(path: Path) -> bool:
    """Give Linux a best-effort hint to evict one prepared input file."""

    if not hasattr(os, "posix_fadvise") or not hasattr(os, "POSIX_FADV_DONTNEED"):
        return False
    with path.open("rb") as cache_handle:
        os.posix_fadvise(cache_handle.fileno(), 0, 0, os.POSIX_FADV_DONTNEED)
    return True


def run_openpyxl(path: Path, expected_cells: int) -> dict[str, int]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=False, data_only=True)
    worksheet = workbook["data"]
    cells = 0
    checksum = 0
    for row in worksheet.iter_rows(values_only=True):
        for value in row:
            cells += 1
            checksum += int(value)
    if cells != expected_cells:
        raise RuntimeError(f"Expected {expected_cells} cells, observed {cells}")
    expected_checksum = expected_cells * (expected_cells - 1) // 2
    if checksum != expected_checksum:
        raise RuntimeError(
            f"Expected checksum {expected_checksum}, observed {checksum}"
        )
    workbook.close()
    return {"cells": cells, "checksum": checksum}


def run_fsspec(path: Path, expected_bytes: int) -> dict[str, int]:
    from fsspec.implementations.local import LocalFileSystem

    # Preparation and measurement are deliberately separate. Ask Linux to
    # evict the prepared file from the page cache where the platform supports
    # it, so the measurement is less likely to describe a warm-cache memcpy.
    cache_hint_applied = drop_file_cache(path)

    filesystem = LocalFileSystem(skip_instance_cache=True)
    total = 0
    reads = 0
    with filesystem.open(str(path), "rb", block_size=MIB) as stream:
        while chunk := stream.read(MIB):
            total += len(chunk)
            reads += 1
    if total != expected_bytes:
        raise RuntimeError(f"Expected {expected_bytes} bytes, observed {total}")
    return {
        "bytes": total,
        "read_calls": reads,
        "cache_hint_applied": cache_hint_applied,
    }


def run_cachetools(size_mib: int) -> dict[str, int]:
    from cachetools import LRUCache

    target_bytes = size_mib * MIB
    chunk_bytes = 256 * 1024
    if target_bytes % chunk_bytes:
        raise ValueError("cachetools size must be divisible by 256 KiB")
    cache = LRUCache(maxsize=target_bytes, getsizeof=len)
    entries = target_bytes // chunk_bytes
    for index in range(entries):
        payload = bytearray(chunk_bytes)
        # Touch every page so RSS reflects committed memory, not only address space.
        for offset in range(0, chunk_bytes, 4096):
            payload[offset] = index % 251
        cache[index] = payload
    if cache.currsize != target_bytes or len(cache) != entries:
        raise RuntimeError(
            f"Expected {entries} entries/{target_bytes} bytes, observed "
            f"{len(cache)} entries/{cache.currsize} bytes"
        )
    checksum = sum(value[0] for value in cache.values())
    return {
        "entries": len(cache),
        "cache_bytes": cache.currsize,
        "checksum": checksum,
    }


def run_smart_open(path: Path, expected_bytes: int) -> dict[str, int | bool]:
    from smart_open import open as smart_open

    cache_hint_applied = drop_file_cache(path)
    total = 0
    reads = 0
    with smart_open(str(path), "rb") as stream:
        while chunk := stream.read(MIB):
            total += len(chunk)
            reads += 1
    if total != expected_bytes:
        raise RuntimeError(f"Expected {expected_bytes} bytes, observed {total}")
    return {
        "bytes": total,
        "read_calls": reads,
        "cache_hint_applied": cache_hint_applied,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=("prepare", "measure"), required=True)
    parser.add_argument(
        "--project",
        choices=("openpyxl", "fsspec", "cachetools", "smart_open"),
        required=True,
    )
    parser.add_argument("--size", type=int, required=True)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    if args.mode == "prepare":
        args.input.parent.mkdir(parents=True, exist_ok=True)
        prepare = {
            "openpyxl": prepare_openpyxl,
            "fsspec": prepare_fsspec,
            "cachetools": prepare_cachetools,
            "smart_open": prepare_fsspec,
        }[args.project]
        details = prepare(args.input, args.size)
        if args.output:
            args.output.write_text(json.dumps(details, indent=2) + "\n")
        return 0

    if args.output is None:
        raise ValueError("--output is required in measure mode")
    if not args.input.is_file():
        raise FileNotFoundError(args.input)

    tracemalloc.start()
    if args.project == "openpyxl":
        result = run_openpyxl(args.input, args.size)
        input_size = args.size
        unit = "cells"
    elif args.project == "fsspec":
        expected_bytes = args.size * MIB
        result = run_fsspec(args.input, expected_bytes)
        input_size = expected_bytes
        unit = "bytes"
    elif args.project == "cachetools":
        result = run_cachetools(args.size)
        input_size = args.size * MIB
        unit = "cache_bytes"
    else:
        expected_bytes = args.size * MIB
        result = run_smart_open(args.input, expected_bytes)
        input_size = expected_bytes
        unit = "bytes"
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    args.output.write_text(
        json.dumps(
            {
                "tracemalloc_peak_bytes": peak,
                "workload": {
                    "project": args.project,
                    "input_size": input_size,
                    "input_unit": unit,
                    "result": result,
                },
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
